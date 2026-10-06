"""DOCX / PPTX / XLSX -> Markdown, with tables, charts, diagrams and pictures kept.

* DOCX  - mammoth (headings, lists, tables) with embedded pictures captured as
          figures in place; text boxes / shapes (which mammoth ignores) appended;
          shape-built diagrams are rendered to images when LibreOffice is present.
* PPTX  - per slide: title, body (with indent levels), tables, chart data, SmartArt
          text, grouped shapes, speaker notes, pictures, and the *connections*
          between shapes recovered from the connector XML ("A -> B") - so a
          box-and-arrow slide survives even with no renderer. Slides that carry
          diagrams are rendered via LibreOffice for the vision model when available.
* XLSX  - one Markdown table per sheet (values), row-capped, plus embedded images.
"""
from __future__ import annotations

import io
import logging
import re
from typing import Any

from . import soffice
from .images import normalise, worth_describing
from .model import IngestError, Limits, RichDocument, marker, md_table, tidy
from .safety import check_zip, parse_xml

log = logging.getLogger("attach.office")

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_DOCX_SHAPE_HINTS = (b"<wps:wsp", b"<wpg:wgp", b"<v:shape", b"<v:group", b"dgm:relIds", b"<w:pict")
_FIG_IMG_RE = re.compile(r"!\[([^\]]*)\]\(fig:(\d+)\)")


# ================================================================== DOCX
def parse_docx(raw: bytes, limits: Limits) -> RichDocument:
    zf = check_zip(raw, limits)
    doc = RichDocument(kind="document")
    try:
        import mammoth
        from markdownify import markdownify
    except ImportError:
        return _docx_fallback(raw, doc)

    pending: dict[int, tuple[bytes, str, str]] = {}

    def _capture(image: Any) -> dict[str, str]:
        with image.open() as fh:
            data = fh.read()
        n = len(pending) + 1
        pending[n] = (data, image.content_type or "", image.alt_text or "")
        return {"src": f"fig:{n}", "alt": image.alt_text or ""}

    try:
        html = mammoth.convert_to_html(io.BytesIO(raw), convert_image=mammoth.images.img_element(_capture)).value
    except Exception as err:  # noqa: BLE001
        log.info("mammoth failed (%s); falling back to python-docx", err)
        return _docx_fallback(raw, doc)
    md = markdownify(html or "", heading_style="ATX", bullets="-")
    doc.bump("tables", len(re.findall(r"(?m)^\| ?---", md)))

    kept = 0
    ids: dict[int, int] = {}
    for n, (data, mime, alt) in sorted(pending.items(), key=lambda kv: -len(kv[1][0])):
        norm = normalise(data, mime)
        if not norm or not worth_describing(norm[2], norm[3], len(norm[0]), limits) or kept >= limits.max_figures * 3:
            continue
        fig = doc.new_figure(norm[0], norm[1], f"picture {n}", caption=alt, width=norm[2], height=norm[3])
        ids[n] = fig.id
        kept += 1
    doc.stats["pictures"] = len(pending)

    def _sub(m: re.Match[str]) -> str:
        fid = ids.get(int(m.group(2)))
        return marker(fid) if fid else ""

    md = _promote_table_headers(_FIG_IMG_RE.sub(_sub, md))
    extra = [t for t in _docx_text_boxes(zf) if t not in md]     # mammoth already inlines some
    if extra:
        md += "\n\n## Text in shapes and text boxes\n\n" + "\n".join(f"- {t}" for t in extra)
        doc.bump("text_boxes", len(extra))
    doc.markdown = tidy(md)
    _docx_render_diagram_pages(raw, zf, doc, limits)
    return doc


_EMPTY_HEADER_RE = re.compile(r"(?m)^\|(?:\s*\|)+\s*\n(\| ?---[^\n]*\n)(\|[^\n]*\n)")


def _promote_table_headers(md: str) -> str:
    """Word tables have no header row, so the HTML->Markdown step emits an empty one.
    Promote the first data row to the header instead."""
    return _EMPTY_HEADER_RE.sub(lambda m: m.group(2) + m.group(1), md)


def _docx_fallback(raw: bytes, doc: RichDocument) -> RichDocument:
    import docx
    d = docx.Document(io.BytesIO(raw))
    parts: list[str] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            parts.append(f"{'#' * min(int(level), 6)} {text}")
        elif style.startswith(("list", "bullet")):
            parts.append(f"- {text}")
        else:
            parts.append(text)
    for table in d.tables:
        parts.append(md_table([[c.text for c in row.cells] for row in table.rows]))
        doc.bump("tables")
    doc.markdown = tidy("\n\n".join(parts))
    return doc


def _docx_text_boxes(zf: Any) -> list[str]:
    try:
        root = parse_xml(zf.read("word/document.xml"))
    except KeyError:
        return []
    seen: list[str] = []
    for box in root.iter(f"{_W}txbxContent"):
        text = " ".join("".join(t.text or "" for t in p.iter(f"{_W}t")).strip() for p in box.iter(f"{_W}p")).strip()
        if text and text not in seen:                      # Word stores each box twice (Choice + Fallback)
            seen.append(text)
    return seen


def _docx_render_diagram_pages(raw: bytes, zf: Any, doc: RichDocument, limits: Limits) -> None:
    """Shape-built diagrams are not pictures: render the pages that carry drawn
    graphics (needs LibreOffice) so the vision model can read them."""
    if not (limits.render_pages and soffice.available()):
        return
    try:
        xml = zf.read("word/document.xml")
    except KeyError:
        return
    if not any(h in xml for h in _DOCX_SHAPE_HINTS):
        return
    pdf_bytes = soffice.convert(raw, "docx", "pdf")
    if not pdf_bytes:
        doc.warnings.append("diagram shapes were found but could not be rendered")
        return
    from . import pdf as pdfmod

    rendered = pdfmod.render_graphic_pages(pdf_bytes, limits, doc, max_pages=max(1, limits.max_figures // 2))
    if rendered:
        doc.markdown += "\n\n## Diagrams drawn in the document\n\n" + "\n\n".join(marker(f) for f in rendered)
        doc.bump("diagram_pages", len(rendered))


# ================================================================== PPTX
def _cell_text(cell: Any) -> str:
    return " ".join(p.text.strip() for p in cell.text_frame.paragraphs if p.text.strip())


def _shape_label(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        return " ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
    return ""


def _arrow_direction(cxn: Any) -> str:
    """'fwd' (begin -> end), 'rev' (end -> begin) or 'none' from the line-end decorations."""
    ln = cxn.find(f".//{_A}ln")
    if ln is None:
        return "fwd"
    head = ln.find(f"{_A}headEnd")
    tail = ln.find(f"{_A}tailEnd")
    has_head = head is not None and head.get("type", "none") != "none"
    has_tail = tail is not None and tail.get("type", "none") != "none"
    if has_tail and not has_head:
        return "fwd"
    if has_head and not has_tail:
        return "rev"
    return "none" if (has_head and has_tail) else "none"


_GROUP, _PICTURE = 6, 13        # MSO_SHAPE_TYPE values


def _stype(shape: Any) -> int | None:
    """shape_type, tolerating the shapes python-pptx cannot classify (it raises for some frames)."""
    try:
        v = shape.shape_type
        return int(v) if v is not None else None
    except Exception:  # noqa: BLE001
        return None


def _walk(shapes: Any):
    """Depth-first over shapes, descending into groups."""
    for sh in shapes:
        yield sh
        if _stype(sh) == _GROUP and getattr(sh, "shapes", None) is not None:
            yield from _walk(sh.shapes)


def _smartart_text(slide: Any, shape: Any) -> list[str]:
    out: list[str] = []
    try:
        for rid in shape._element.xpath(".//*[local-name()='relIds']/@*[local-name()='dm']"):
            part = slide.part.related_part(rid)
            root = parse_xml(part.blob)
            out += [t.text.strip() for t in root.iter(f"{_A}t") if (t.text or "").strip()]
    except Exception:  # noqa: BLE001
        pass
    return out


def _chart_md(shape: Any) -> str:
    try:
        chart = shape.chart
        title = chart.chart_title.text_frame.text.strip() if chart.has_title and chart.chart_title.has_text_frame else ""
        lines = [f"Chart ({str(chart.chart_type).split('.')[-1].split(' ')[0].lower()})" + (f": {title}" if title else "")]
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        rows = [["Series", *cats]]
        for s in plot.series:
            rows.append([s.name or "", *["" if v is None else f"{v:g}" for v in s.values]])
        return "\n".join(lines) + "\n\n" + md_table(rows)
    except Exception:  # noqa: BLE001 - exotic chart
        return "Chart (data could not be read)"


class _Slide:
    def __init__(self, index: int) -> None:
        self.index = index
        self.title = ""
        self.body: list[str] = []
        self.tables: list[str] = []
        self.charts: list[str] = []
        self.smartart: list[str] = []
        self.notes = ""
        self.edges: list[str] = []
        self.pictures: list[tuple[bytes, str, str]] = []   # blob, ext, alt
        self.connectors = 0
        self.text_shapes = 0
        self.groups = 0
        self.hidden = False

    @property
    def diagram_score(self) -> float:
        s = len(self.pictures) * 3 + len(self.charts) * 3 + (4 if self.smartart else 0) \
            + self.connectors * 2 + self.groups
        if self.text_shapes >= 5:
            s += 2
        return float(s)


def _read_slide(slide: Any, index: int) -> _Slide:
    out = _Slide(index)
    out.hidden = slide._element.get("show") == "0"
    try:
        if slide.shapes.title is not None:
            out.title = _shape_label(slide.shapes.title)
    except Exception:  # noqa: BLE001
        pass
    labels: dict[str, str] = {}
    cxns: list[Any] = []
    ordered = sorted(_walk(slide.shapes), key=lambda s: (round((s.top or 0) / 228600), s.left or 0))
    for sh in ordered:
        if _stype(sh) == _GROUP:
            out.groups += 1
            continue
        if sh.element.tag == f"{_P}cxnSp":
            cxns.append(sh.element)
            out.connectors += 1
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = [[_cell_text(c) for c in r.cells] for r in sh.table.rows]
            md = md_table(rows)
            if md:
                out.tables.append(md)
            continue
        if getattr(sh, "has_chart", False) and sh.has_chart:
            out.charts.append(_chart_md(sh))
            continue
        if _stype(sh) == _PICTURE or sh.element.tag == f"{_P}pic":
            try:
                alt = sh.element.xpath(".//*[local-name()='cNvPr']/@descr")
                out.pictures.append((sh.image.blob, sh.image.ext, alt[0] if alt else ""))
            except Exception:  # noqa: BLE001
                pass
            continue
        if sh.element.tag == f"{_P}graphicFrame" and not getattr(sh, "has_table", False):
            sa = _smartart_text(slide, sh)
            if sa:
                out.smartart += sa
                continue
        text = _shape_label(sh)
        if not text:
            continue
        sid = str(sh.shape_id)
        labels[sid] = text
        is_title = out.title and sh.is_placeholder and text == out.title
        if is_title:
            continue
        out.text_shapes += 1
        lines = []
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip() or p.text.strip()
            if t:
                lines.append(("  " * min(p.level, 4)) + f"- {t}")
        out.body += lines
    for c in cxns:
        cnv = c.find(f".//{_P}cNvCxnSpPr")
        if cnv is None:
            continue
        st, en = cnv.find(f"{_A}stCxn"), cnv.find(f"{_A}endCxn")
        a = labels.get(st.get("id", "")) if st is not None else None
        b = labels.get(en.get("id", "")) if en is not None else None
        if a and b:
            arrow = {"fwd": "→", "rev": "←", "none": "—"}[_arrow_direction(c)]
            out.edges.append(f"{a} {arrow} {b}")
    try:
        if slide.has_notes_slide:
            out.notes = slide.notes_slide.notes_text_frame.text.strip()
    except Exception:  # noqa: BLE001
        pass
    return out


def parse_pptx(raw: bytes, limits: Limits) -> RichDocument:
    check_zip(raw, limits)
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(raw))
    except ImportError as err:  # pragma: no cover
        raise IngestError("PowerPoint support is not installed on this server") from err
    except Exception as err:  # noqa: BLE001
        raise IngestError(f"could not open the presentation ({type(err).__name__})") from err

    doc = RichDocument(kind="document")
    slides = list(prs.slides)
    doc.stats["slides"] = len(slides)
    read: list[_Slide] = []
    for i, s in enumerate(slides[: limits.max_pages], start=1):
        try:
            read.append(_read_slide(s, i))
        except Exception as err:  # noqa: BLE001 - one odd slide must not lose the deck
            log.warning("slide %s unreadable: %s", i, err)
            doc.warnings.append(f"slide {i} could not be read")
            read.append(_Slide(i))
    if len(slides) > len(read):
        doc.warnings.append(f"only the first {len(read)} of {len(slides)} slides were read (limit)")

    rendered = _render_slides(raw, read, len(slides), doc, limits)
    pictures_budget = limits.max_figures * 3
    out: list[str] = []
    for sl in read:
        head = f"## Slide {sl.index}" + (f": {sl.title}" if sl.title else "")
        if sl.hidden:
            head += " (hidden)"
        out.append(head)
        out += sl.body
        if sl.smartart:
            out.append("SmartArt diagram text: " + " · ".join(sl.smartart))
        for t in sl.tables:
            out.append("\n" + t + "\n")
            doc.bump("tables")
        for c in sl.charts:
            out.append("\n" + c + "\n")
            doc.bump("charts")
        if sl.edges:
            out.append("Diagram connections:\n" + "\n".join(f"- {e}" for e in sl.edges))
            doc.bump("diagram_edges", len(sl.edges))
        fid = rendered.get(sl.index)
        if fid:
            out.append(marker(fid))
        else:
            for blob, ext, alt in sl.pictures:
                norm = normalise(blob, ext)
                if not norm:
                    if alt:
                        out.append(f"[picture: {alt}]")
                    continue
                if not worth_describing(norm[2], norm[3], len(norm[0]), limits) or pictures_budget <= 0:
                    if alt:
                        out.append(f"[picture: {alt}]")
                    continue
                pictures_budget -= 1
                fig = doc.new_figure(norm[0], norm[1], f"slide {sl.index}", caption=alt,
                                     width=norm[2], height=norm[3], context=sl.title)
                out.append(marker(fig.id))
        if sl.notes:
            out.append(f"Speaker notes: {sl.notes}")
        out.append("")
    doc.markdown = tidy("\n".join(out))
    return doc


def _render_slides(raw: bytes, slides: list[_Slide], total: int, doc: RichDocument, limits: Limits) -> dict[int, int]:
    """Render diagram-bearing slides to images (LibreOffice -> PDF -> pdfium)."""
    wanted = sorted((s for s in slides if s.diagram_score >= 3), key=lambda s: (-s.diagram_score, s.index))
    wanted = wanted[: limits.max_figures]
    if not wanted or not limits.render_pages or not soffice.available():
        if wanted and limits.render_pages:
            doc.warnings.append("slides with diagrams were read as text and pictures only "
                                "(slide rendering needs LibreOffice, which is not installed)")
        return {}
    pdf_bytes = soffice.convert(raw, "pptx", "pdf")
    if not pdf_bytes:
        doc.warnings.append("slides could not be rendered; they were read as text and pictures only")
        return {}
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(pdf_bytes)
    except Exception:  # noqa: BLE001
        doc.warnings.append("slide rendering is unavailable")
        return {}
    try:
        visible = [s.index for s in slides if not s.hidden]
        if len(pdf) == total:
            page_of = {s.index: s.index - 1 for s in slides}
        elif len(pdf) == len(visible) and len(slides) == total:
            page_of = {idx: n for n, idx in enumerate(visible)}
        else:
            doc.warnings.append("slide numbering did not match the rendered output; slides were read as text only")
            return {}
        out: dict[int, int] = {}
        for sl in sorted(wanted, key=lambda s: s.index):
            n = page_of.get(sl.index)
            if n is None:
                continue
            try:
                page = pdf[n]
                w, h = page.get_size()
                scale = max(0.5, min(2.5, limits.figure_max_edge / max(w, h)))
                img = page.render(scale=scale).to_pil().convert("RGB")
                buf = io.BytesIO()
                img.save(buf, format="JPEG", quality=85)
                fig = doc.new_figure(buf.getvalue(), "image/jpeg", f"slide {sl.index}", kind="slide",
                                     width=img.width, height=img.height, context=sl.title,
                                     caption=("diagram slide" if sl.connectors or sl.groups else ""))
                out[sl.index] = fig.id
            except Exception as err:  # noqa: BLE001
                log.warning("render slide %s failed: %s", sl.index, err)
        return out
    finally:
        pdf.close()


# ================================================================== XLSX
def parse_xlsx(raw: bytes, limits: Limits) -> RichDocument:
    zf = check_zip(raw, limits)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except ImportError as err:  # pragma: no cover
        raise IngestError("spreadsheet support is not installed on this server") from err
    except Exception as err:  # noqa: BLE001
        raise IngestError(f"could not open the spreadsheet ({type(err).__name__})") from err

    doc = RichDocument(kind="document")
    out: list[str] = []
    doc.stats["sheets"] = len(wb.worksheets)
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        truncated = False
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                rows.append(cells)
            if len(rows) > limits.max_rows_per_sheet:
                truncated = True
                break
        if not rows:
            continue
        out.append(f"## {ws.title}")
        out.append(md_table(rows[: limits.max_rows_per_sheet]))
        if truncated:
            out.append(f"_(sheet truncated to the first {limits.max_rows_per_sheet} rows)_")
            doc.warnings.append(f"sheet '{ws.title}' was truncated to {limits.max_rows_per_sheet} rows")
        out.append("")
        doc.bump("tables")
    wb.close()
    kept = 0
    for name in sorted(n for n in zf.namelist() if n.startswith("xl/media/")):
        norm = normalise(zf.read(name), name.rsplit(".", 1)[-1])
        if norm and worth_describing(norm[2], norm[3], len(norm[0]), limits) and kept < limits.max_figures:
            fig = doc.new_figure(norm[0], norm[1], f"workbook image {kept + 1}", width=norm[2], height=norm[3])
            out.append(marker(fig.id))
            kept += 1
    doc.markdown = tidy("\n".join(out))
    return doc

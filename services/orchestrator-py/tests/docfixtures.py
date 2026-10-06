"""Builders for real PDF / DOCX / PPTX / draw.io / image fixtures used by the
document-ingestion tests (generated on the fly - no binary files in the repo)."""
from __future__ import annotations

import io

from PIL import Image, ImageDraw


def png(w: int = 320, h: int = 200, text: str = "Order Service", color=(30, 90, 200)) -> bytes:
    """A diagram-like picture: boxes, an arrow and a label (big enough to pass the icon filter)."""
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([10, 10, w // 2 - 10, h // 2], outline=color, width=3)
    d.rectangle([w // 2 + 10, h // 2, w - 10, h - 10], outline=color, width=3)
    d.line([w // 2 - 10, h // 4, w // 2 + 10, h * 3 // 4], fill=color, width=3)
    d.text((20, 20), text, fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def tiny_png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (16, 16), "red").save(buf, format="PNG")
    return buf.getvalue()


def make_pdf(*, pages: int = 1, table: bool = True, diagram: bool = True, picture: bool = False,
             scanned: bool = False, header: bool = False) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Table, TableStyle

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    for p in range(1, pages + 1):
        if header:
            c.setFont("Helvetica", 8)
            c.drawString(40, height - 24, "ACME Confidential Requirements")
            c.drawString(width / 2, 20, f"Page {p}")
        if scanned:
            c.drawImage(ImageReader(io.BytesIO(png(600, 800, "scanned"))), 40, 100, 500, 650)
            c.showPage()
            continue
        c.setFont("Helvetica-Bold", 22)
        c.drawString(50, height - 80, f"{p}. Payment Requirements")
        c.setFont("Helvetica", 11)
        y = height - 110
        for line in ("The system shall let customers pay by card.", "Refunds must complete within 5 days."):
            c.drawString(50, y, line)
            y -= 16
        c.setFont("Helvetica-Bold", 14)
        c.drawString(50, y - 10, "Acceptance criteria")
        c.setFont("Helvetica", 11)
        c.drawString(50, y - 30, "- Card declined shows an error")
        y -= 60
        if table:
            t = Table([["ID", "Requirement", "Priority"], ["R1", "Pay by card", "High"], ["R2", "Refund", "Medium"]])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
            tw, th = t.wrapOn(c, width, height)
            t.drawOn(c, 50, y - th)
            y -= th + 20
        if diagram:  # pure vector boxes + arrows: not an embedded image
            for i in range(5):
                c.rect(60 + i * 90, y - 80, 70, 40)
                c.line(130 + i * 90, y - 60, 150 + i * 90, y - 60)
            y -= 110
        if picture:
            c.drawImage(ImageReader(io.BytesIO(png())), 60, y - 160, 240, 150)
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(*, picture: bool = True, table: bool = True, text_box: bool = False, shapes: bool = False) -> bytes:
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.add_heading("Order Management Requirements", level=1)
    d.add_paragraph("The platform shall let customers place and track orders.")
    d.add_heading("Functional requirements", level=2)
    d.add_paragraph("Place an order", style="List Bullet")
    d.add_paragraph("Cancel an order", style="List Bullet")
    if table:
        t = d.add_table(rows=3, cols=3)
        for r, row in enumerate([["ID", "Requirement", "Priority"], ["R1", "Place order", "High"], ["R2", "Cancel", "Low"]]):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    if picture:
        d.add_heading("Architecture", level=2)
        d.add_picture(io.BytesIO(png(400, 260, "Context diagram")), width=Inches(4))
    d.add_paragraph("End of document.")
    buf = io.BytesIO()
    d.save(buf)
    raw = buf.getvalue()
    if text_box or shapes:
        raw = _inject_docx_xml(raw, text_box=text_box, shapes=shapes)
    return raw


def _inject_docx_xml(raw: bytes, *, text_box: bool, shapes: bool) -> bytes:
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                box = ""
                if text_box:
                    box = ('<w:p><w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox>'
                           '<w:txbxContent><w:p><w:r><w:t>Callout: PCI scope ends at the gateway</w:t></w:r></w:p>'
                           '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>')
                xml = xml.replace("</w:body>", box + "</w:body>")
                data = xml.encode("utf-8")
            dst.writestr(item, data)
    return out.getvalue()


def make_pptx(*, picture: bool = True, table: bool = True, chart: bool = True, diagram: bool = True,
              notes: bool = True) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Scope and goals"
    s.placeholders[1].text_frame.text = "Launch checkout in Q3"
    p = s.placeholders[1].text_frame.add_paragraph()
    p.text = "Support card and wallet payments"
    p.level = 1
    if notes:
        s.notes_slide.notes_text_frame.text = "Mention the PCI audit date."

    if table or chart or picture:
        s2 = prs.slides.add_slide(prs.slide_layouts[5])
        s2.shapes.title.text = "Data and targets"
        if table:
            gf = s2.shapes.add_table(3, 2, Inches(0.5), Inches(1.5), Inches(4), Inches(1.2))
            for r, row in enumerate([["Metric", "Target"], ["Latency", "200ms"], ["Uptime", "99.9%"]]):
                for c, v in enumerate(row):
                    gf.table.cell(r, c).text = v
        if chart:
            cd = CategoryChartData()
            cd.categories = ["Q1", "Q2", "Q3"]
            cd.add_series("Orders", (120, 250, 400))
            s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1.5), Inches(4), Inches(3), cd)
        if picture:
            s2.shapes.add_picture(io.BytesIO(png(400, 260, "Wireframe")), Inches(0.5), Inches(3.5), Inches(3), Inches(2))

    if diagram:
        s3 = prs.slides.add_slide(prs.slide_layouts[5])
        s3.shapes.title.text = "Solution overview"
        a = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2), Inches(2), Inches(1))
        a.text_frame.text = "Web App"
        b = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(2), Inches(2), Inches(1))
        b.text_frame.text = "Order API"
        c = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.5), Inches(2), Inches(2), Inches(1))
        c.text_frame.text = "Payments DB"
        for left, right in ((a, b), (b, c)):
            cx = s3.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
            cx.begin_connect(left, 3)
            cx.end_connect(right, 1)
            ln = cx.line._get_or_add_ln()
            from lxml import etree
            tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
            tail.set("type", "triangle")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_xlsx(*, rows: int = 4) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Backlog"
    ws.append(["ID", "Story", "Points"])
    for i in range(1, rows):
        ws.append([f"S{i}", f"Story {i}", i * 2])
    wb.create_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DRAWIO = """<mxfile><diagram name="Flow"><mxGraphModel><root>
<mxCell id="0"/><mxCell id="1" parent="0"/>
<mxCell id="2" value="Client" vertex="1" parent="1"/>
<mxCell id="3" value="API &lt;b&gt;Gateway&lt;/b&gt;" vertex="1" parent="1"/>
<mxCell id="4" value="Orders DB" vertex="1" parent="1"/>
<mxCell id="5" value="HTTPS" edge="1" source="2" target="3" parent="1"/>
<mxCell id="6" value="" edge="1" source="3" target="4" parent="1"/>
</root></mxGraphModel></diagram></mxfile>"""


def make_pptx_with_smartart_and_group() -> bytes:
    """A slide holding a SmartArt frame (text lives in a diagram-data part) and a group of two boxes."""
    from lxml import etree
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Process"
    grp = slide.shapes.add_group_shape()
    for i, label in enumerate(("Intake", "Review")):
        box = grp.shapes.add_shape(1, Inches(0.5 + i * 3), Inches(4), Inches(2), Inches(0.8))
        box.text_frame.text = label
    data = (b'<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
            b'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst>'
            b'<dgm:pt modelId="1"><dgm:t><a:bodyPr/><a:p><a:r><a:t>Collect</a:t></a:r></a:p></dgm:t></dgm:pt>'
            b'<dgm:pt modelId="2"><dgm:t><a:bodyPr/><a:p><a:r><a:t>Approve</a:t></a:r></a:p></dgm:t></dgm:pt>'
            b'</dgm:ptLst></dgm:dataModel>')
    part = Part(PackURI("/ppt/diagrams/data1.xml"),
                "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml", prs.part.package, data)
    rid = slide.part.relate_to(part, RT.DIAGRAM_DATA)
    frame = etree.fromstring(
        '<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<p:nvGraphicFramePr><p:cNvPr id="900" name="SmartArt"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        '<p:xfrm><a:off x="457200" y="1371600"/><a:ext cx="4572000" cy="2000000"/></p:xfrm>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
        f'<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" r:dm="{rid}" r:lo="rId1" r:qs="rId1" r:cs="rId1"/>'
        '</a:graphicData></a:graphic></p:graphicFrame>')
    slide.shapes._spTree.append(frame)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

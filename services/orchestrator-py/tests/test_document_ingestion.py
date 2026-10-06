"""Rich document ingestion: text, tables, pictures, diagrams and slides from real
PDF / DOCX / PPTX / XLSX / draw.io files (generated on the fly), the figure
pipeline, budget fitting, and the upload route."""
from __future__ import annotations

import io
import zipfile
from types import SimpleNamespace

import pytest

from app.services import attachment_extract as ax
from app.services.documents import Limits, allocate, fit_document, ingest, parse, soffice
from app.services.documents.diagrams import parse_drawio, parse_svg, parse_vsdx
from app.services.documents.figures import apply_descriptions, describe_figures, strip_markers
from app.services.documents.model import FIG_RE, IngestError, RichDocument, marker
from app.services.documents.office import parse_docx, parse_pptx, parse_xlsx
from app.services.documents.pdf import parse_pdf

from .docfixtures import (DRAWIO, make_docx, make_pdf, make_pptx, make_pptx_with_smartart_and_group, make_xlsx,
                          png, tiny_png)


def _renders_slides() -> bool:
    return soffice.available() and soffice.convert(make_pptx(), "pptx", "pdf") is not None


needs_soffice = pytest.mark.skipif(not _renders_slides(), reason="LibreOffice (Impress) not installed")


async def _describe(fig):
    return f"Diagram of {fig.label}: Web App -> Order API.", "vision:test"


# ================================================================== PDF
def test_pdf_keeps_headings_lists_and_real_tables():
    doc = parse_pdf(make_pdf(table=True, diagram=False), Limits())
    md = doc.markdown
    assert "# 1. Payment Requirements" in md
    assert "### Acceptance criteria" in md and "- Card declined shows an error" in md
    assert "| ID | Requirement | Priority |" in md and "| R1 | Pay by card | High |" in md
    assert "The system shall let customers pay by card. Refunds must complete within 5 days." in md  # one paragraph
    assert doc.stats["pages"] == 1 and doc.stats["tables"] == 1 and not doc.figures


def test_pdf_drops_running_headers_and_footers_but_keeps_headings():
    doc = parse_pdf(make_pdf(pages=6, header=True, diagram=False), Limits())
    assert "ACME Confidential" not in doc.markdown
    assert doc.markdown.count("Payment Requirements") == 6


def test_pdf_vector_diagram_page_is_rendered_for_the_vision_model():
    doc = parse_pdf(make_pdf(diagram=True), Limits())
    (fig,) = doc.figures
    assert fig.kind == "page" and fig.label == "page 1" and fig.mime == "image/jpeg"
    assert max(fig.width, fig.height) <= 1568 and len(fig.data) > 5_000
    assert marker(fig.id) in doc.markdown


def test_pdf_page_with_picture_is_rendered_and_scanned_page_is_prioritised():
    assert parse_pdf(make_pdf(table=False, diagram=False, picture=True), Limits()).figures
    scanned = parse_pdf(make_pdf(scanned=True, table=False, diagram=False), Limits())
    (fig,) = scanned.figures
    assert fig.caption == "scanned page"


def test_pdf_figure_limit_is_reported_not_silent():
    doc = parse_pdf(make_pdf(pages=6), Limits(max_figures=2))
    assert len(doc.figures) == 2
    assert any("not analysed (figure limit 2)" in w for w in doc.warnings)


def test_pdf_page_limit_is_reported():
    doc = parse_pdf(make_pdf(pages=5, diagram=False), Limits(max_pages=2))
    assert any("only the first 2 of 5 pages" in w for w in doc.warnings)
    assert "<!-- page 3 -->" not in doc.markdown


def test_pdf_time_budget_is_reported():
    doc = parse_pdf(make_pdf(pages=3, diagram=False), Limits(max_seconds=0))
    assert any("time budget reached" in w for w in doc.warnings)


def test_corrupt_pdf_is_refused_with_a_reason():
    with pytest.raises(IngestError, match="corrupt or password-protected"):
        parse_pdf(b"%PDF-1.4 not really", Limits())


# ================================================================== DOCX
def test_docx_structure_tables_and_picture_in_place():
    doc = parse_docx(make_docx(), Limits())
    md = doc.markdown
    assert "# Order Management Requirements" in md and "## Functional requirements" in md
    assert "- Place an order" in md and "- Cancel an order" in md
    # Word tables have no header row: the first row must become the header, not an empty one
    assert "| ID | Requirement | Priority |\n| --- | --- | --- |\n| R1 | Place order | High |" in md
    (fig,) = doc.figures
    assert md.index("## Architecture") < md.index(marker(fig.id)) < md.index("End of document.")
    assert fig.width == 400 and doc.stats["tables"] == 1


def test_docx_text_boxes_are_captured_once():
    md = parse_docx(make_docx(text_box=True, picture=False), Limits()).markdown
    assert md.count("PCI scope ends at the gateway") == 1


def test_docx_icons_are_not_sent_to_vision():
    import docx

    d = docx.Document()
    d.add_paragraph("Icon below")
    d.add_picture(io.BytesIO(tiny_png()))
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_docx(buf.getvalue(), Limits())
    assert doc.figures == [] and doc.stats["pictures"] == 1


def test_not_a_zip_is_refused():
    with pytest.raises(IngestError, match="not a valid Office"):
        parse_docx(b"this is not a docx", Limits())


def test_archive_bombs_are_refused():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<a/>" * 50_000)
    with pytest.raises(IngestError, match="unsafe size"):
        parse_docx(buf.getvalue(), Limits(max_unzipped_bytes=1000))


# ================================================================== PPTX
def test_pptx_text_tables_charts_notes_and_diagram_connections():
    doc = parse_pptx(make_pptx(), Limits(render_pages=False))
    md = doc.markdown
    assert "## Slide 1: Scope and goals" in md
    assert "- Launch checkout in Q3" in md and "  - Support card and wallet payments" in md
    assert "Speaker notes: Mention the PCI audit date." in md
    assert "| Metric | Target |" in md and "| Latency | 200ms |" in md
    assert "| Series | Q1 | Q2 | Q3 |" in md and "| Orders | 120 | 250 | 400 |" in md
    # box-and-arrow slide: the connections are recovered from the connector XML
    assert "- Web App → Order API" in md and "- Order API → Payments DB" in md
    assert doc.stats["slides"] == 3 and doc.stats["diagram_edges"] == 2 and doc.stats["charts"] == 1
    assert [f.label for f in doc.figures] == ["slide 2"]            # the wireframe picture
    assert any("LibreOffice" in w for w in doc.warnings) or soffice.available()


@needs_soffice
def test_pptx_diagram_slides_are_rendered_when_libreoffice_exists():
    doc = parse_pptx(make_pptx(), Limits())
    kinds = {f.label: f.kind for f in doc.figures}
    assert kinds == {"slide 2": "slide", "slide 3": "slide"}
    # the rendered slide replaces the separate picture for that slide (no duplicate)
    assert len(doc.figures) == 2


def test_pptx_smartart_text_and_grouped_shapes_are_read_without_crashing():
    doc = parse_pptx(make_pptx_with_smartart_and_group(), Limits(render_pages=False))
    md = doc.markdown
    assert "SmartArt diagram text: Collect · Approve" in md
    assert "- Intake" in md and "- Review" in md                      # shapes inside the group
    assert doc.warnings == [] or all("could not be read" not in w for w in doc.warnings)


def test_pptx_hidden_and_corrupt_inputs():
    with pytest.raises(IngestError):
        parse_pptx(b"nope", Limits())


# ================================================================== XLSX
def test_xlsx_tables_and_row_cap():
    doc = parse_xlsx(make_xlsx(rows=50), Limits(max_rows_per_sheet=10))
    assert "## Backlog" in doc.markdown and "| ID | Story | Points |" in doc.markdown
    assert "## Empty" not in doc.markdown
    assert any("truncated to 10 rows" in w for w in doc.warnings)


# ================================================================== diagrams
def test_drawio_components_and_directed_connections():
    md = parse_drawio(DRAWIO.encode(), Limits()).markdown
    assert "- API Gateway" in md                      # html label stripped
    assert "- Client → API Gateway [HTTPS]" in md and "- API Gateway → Orders DB" in md


def test_drawio_compressed_diagram():
    import base64
    import urllib.parse
    import zlib

    inner = DRAWIO.split('<diagram name="Flow">')[1].split("</diagram>")[0]
    packed = base64.b64encode(zlib.compress(urllib.parse.quote(inner).encode())[2:-4]).decode()
    md = parse_drawio(f'<mxfile><diagram name="C">{packed}</diagram></mxfile>'.encode(), Limits()).markdown
    assert "Client → API Gateway [HTTPS]" in md


def test_svg_text_extraction_and_xxe_is_not_expanded():
    svg = (b'<?xml version="1.0"?><!DOCTYPE s [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
           b'<svg xmlns="http://www.w3.org/2000/svg"><title>Flow</title><text>Start &x;</text></svg>')
    md = parse_svg(svg, Limits()).markdown
    assert "Title: Flow" in md and "root:" not in md


def test_vsdx_shapes_and_connections():
    page = ('<PageContents xmlns="http://schemas.microsoft.com/office/visio/2012/main"><Shapes>'
            '<Shape ID="1"><Text>Client</Text></Shape><Shape ID="2"><Text>Server</Text></Shape><Shape ID="3"/></Shapes>'
            '<Connects><Connect FromSheet="3" FromCell="BeginX" ToSheet="1"/>'
            '<Connect FromSheet="3" FromCell="EndX" ToSheet="2"/></Connects></PageContents>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("visio/pages/page1.xml", page)
    assert "- Client → Server" in parse_vsdx(buf.getvalue(), Limits()).markdown


# ================================================================== dispatch
def test_dispatch_by_extension_sniffing_and_fallbacks():
    assert parse(b"a,b\n1,2", "t.csv", "text/csv", Limits()).kind == "text"
    assert parse(b"# Spec\nMUST do X", "noext", "", Limits()).markdown.startswith("# Spec")
    assert "Title" in parse(b"<html><body><h1>Title</h1><script>x()</script></body></html>", "a.html", "", Limits()).markdown
    assert parse(make_docx(), "renamed.bin", "application/octet-stream", Limits()).stats["tables"] == 1  # sniffed
    assert parse(DRAWIO.encode(), "diagram.xml", "application/xml", Limits()).kind == "diagram"
    binary = parse(bytes(range(256)) * 8, "blob.bin", "application/octet-stream", Limits())
    assert binary.kind == "binary" and "not inlined" in binary.warnings[0]


def test_windows_1252_text_is_decoded():
    assert "café" in parse("café".encode("cp1252"), "notes.txt", "", Limits()).markdown


def test_legacy_formats_need_libreoffice_and_say_so(monkeypatch):
    monkeypatch.setattr(soffice, "available", lambda: False)
    with pytest.raises(IngestError, match="save the file as .pptx"):
        parse(b"\xd0\xcf\x11\xe0", "old.ppt", "", Limits())


def test_legacy_conversion_roundtrip(monkeypatch):
    monkeypatch.setattr(soffice, "available", lambda: True)
    monkeypatch.setattr(soffice, "convert", lambda raw, src, target: make_docx(picture=False))
    doc = parse(b"\xd0\xcf\x11\xe0", "old.doc", "", Limits())
    assert "Order Management Requirements" in doc.markdown and doc.warnings[0].startswith("converted from .doc")


# ================================================================== figures
def _doc_with_figures(n: int, dup: bool = False) -> RichDocument:
    doc = RichDocument(kind="document")
    parts = []
    for i in range(n):
        fig = doc.new_figure(png(300, 200, "same" if dup else f"fig{i}"), "image/png", f"page {i + 1}",
                             width=300, height=200 + i)
        parts.append(f"Text {i}\n\n{marker(fig.id)}")
    doc.markdown = "\n\n".join(parts)
    return doc


async def test_figure_descriptions_land_where_the_figure_was():
    doc = _doc_with_figures(2)
    apply_descriptions(doc, await describe_figures(doc, _describe, Limits()))
    md = doc.markdown
    assert not FIG_RE.search(md)
    assert md.index("Text 0") < md.index("Diagram of page 1") < md.index("Text 1") < md.index("Diagram of page 2")
    assert "> **Figure — page 1** · _read by vision:test_" in md
    assert doc.stats["figures_found"] == 2 and doc.stats["figures_described"] == 2


async def test_identical_figures_are_described_once():
    doc = _doc_with_figures(4, dup=True)
    calls = []

    async def counting(fig):
        calls.append(fig.id)
        return "same diagram", "vision:test"

    apply_descriptions(doc, await describe_figures(doc, counting, Limits()))
    assert len(calls) == 1 and doc.markdown.count("same diagram") == 4


async def test_figure_cap_and_unreadable_figures_are_reported():
    doc = _doc_with_figures(5)
    apply_descriptions(doc, await describe_figures(doc, _describe, Limits(max_figures=2)))
    assert doc.stats["figures_described"] == 2 and doc.stats["figures_unread"] == 3
    assert doc.markdown.count("not analysed — figure limit reached (2 per document)") == 3


async def test_a_failing_or_unavailable_describer_never_fails_the_document():
    async def boom(fig):
        raise RuntimeError("gateway down")

    async def none(fig):
        return None

    for describer in (boom, none):
        doc = _doc_with_figures(1)
        apply_descriptions(doc, await describe_figures(doc, describer, Limits()))
        assert "not analysed — no vision model or OCR was available" in doc.markdown and "Text 0" in doc.markdown


async def test_figure_time_budget_skips_the_rest():
    import asyncio

    async def slow(fig):
        await asyncio.sleep(5)
        return "late", "vision:test"

    doc = _doc_with_figures(2)
    apply_descriptions(doc, await describe_figures(doc, slow, Limits(figure_seconds=1.2, figure_concurrency=1)))
    assert "time budget" in doc.markdown and "late" not in doc.markdown


def test_strip_markers_leaves_an_honest_note():
    doc = _doc_with_figures(1)
    strip_markers(doc)
    assert "[Figure — page 1: not analysed]" in doc.markdown and not FIG_RE.search(doc.markdown)


async def test_ingest_end_to_end_merges_descriptions_and_bounds_size():
    doc = await ingest(make_pptx(), "deck.pptx", "", limits=Limits(render_pages=False), describer=_describe)
    assert "Diagram of slide 2" in doc.markdown and doc.stats["figures_described"] == 1
    big = await ingest(b"x " * 5000, "big.txt", "", limits=Limits(max_chars=1000))
    assert len(big.markdown) == 1000 and any("cut at 1,000" in w for w in big.warnings)


# ================================================================== extract_rich (vision wiring)
class _Llm:
    def __init__(self, provider="bedrock"):
        self.provider, self.prompts = provider, []

    async def generate(self, *, intent, messages, **kw):
        self.prompts.append(messages[0]["content"])
        return SimpleNamespace(provider=self.provider, content="A context diagram with Web App and Order API.")


async def test_extract_rich_uses_the_vision_model_with_document_context():
    llm = _Llm()
    res = await ax.extract_rich(make_docx(), "reqs.docx", "", llm=llm, settings=SimpleNamespace(ATTACHMENT_VISION="auto"))
    assert "A context diagram with Web App and Order API." in res.text
    assert res.stats["figures_described"] == 1 and res.method == "vision:bedrock" and res.kind == ax.KIND_DOCUMENT
    assert '"reqs.docx"' in llm.prompts[0] and "picture 1" in llm.prompts[0]      # figure prompt carries its location
    assert res.outline and res.outline[0] == "# Order Management Requirements"


async def test_extract_rich_standalone_image_uses_the_image_prompt():
    llm = _Llm()
    res = await ax.extract_rich(png(), "whiteboard.png", "image/png", llm=llm, settings=None)
    assert res.kind == ax.KIND_IMAGE and "context diagram" in res.text and "Transcribe ALL text" in llm.prompts[0]


async def test_extract_rich_degrades_to_structure_when_no_real_vision_provider():
    res = await ax.extract_rich(make_pptx(), "deck.pptx", "", llm=_Llm("mock"),
                                settings=SimpleNamespace(ATTACHMENT_VISION="auto", ATTACHMENT_RENDER_PAGES="off"))
    assert "- Web App → Order API" in res.text                       # diagram structure survives with no model at all
    assert res.stats["figures_described"] == 0 and res.warnings and "could not be analysed" in res.note


async def test_extract_rich_ocr_only_mode_never_calls_the_llm():
    llm = _Llm()
    await ax.extract_rich(make_docx(), "reqs.docx", "", llm=llm, settings=SimpleNamespace(ATTACHMENT_VISION="ocr"))
    assert llm.prompts == []


async def test_extract_rich_refuses_corrupt_files():
    with pytest.raises(ax.IngestError):
        await ax.extract_rich(b"junk", "x.pdf", "application/pdf", llm=None)


# ================================================================== fitting a document into the prompt
def _long_doc() -> str:
    parts = ["# Platform requirements", "Overview of the platform and its goals."]
    for i in range(1, 21):
        body = f"Section {i} body. " + ("Filler sentence about nothing in particular. " * 120)
        if i == 13:
            body = ("Refund processing rules: refunds over 500 EUR need dual approval. " * 3) + body
        parts += [f"## Section {i}: {'Refunds' if i == 13 else 'Topic ' + str(i)}", body]
    return "\n\n".join(parts)


def test_small_documents_are_untouched():
    assert fit_document("# A\nshort", 1000, "anything") == "# A\nshort"


def test_large_documents_keep_every_section_and_favour_relevant_ones():
    doc = _long_doc()
    out = fit_document(doc, 9_000, "design the refund approval workflow")
    assert len(out) <= 9_000 + 400 and len(out) < len(doc) / 4
    for i in range(1, 21):                                          # every section keeps its heading AND its opening text
        assert f"## Section {i}:" in out and f"Section {i} body." in out, i
    assert out.count("Refund processing rules") == 3               # the relevant section is kept whole
    assert "rest of this section omitted" in out and "Document outline:" in out
    assert out.index("Section 1:") < out.index("Section 13:") < out.index("Section 20:")   # original order


def test_a_tiny_budget_marks_whole_omitted_sections_instead_of_hiding_them():
    out = fit_document(_long_doc(), 1_200, "refund")
    assert "omitted to fit the prompt budget" in out and len(out) <= 1_200
    assert "Document outline:" in out                              # the outline still names every section


def test_fitting_is_deterministic_and_never_cuts_a_table_row():
    table = "\n".join(["| ID | Requirement |", "| --- | --- |", *[f"| R{i} | requirement number {i} |" for i in range(200)]])
    doc = "# T\n\n" + table
    a, b = fit_document(doc, 1500, ""), fit_document(doc, 1500, "")
    assert a == b
    for line in a.splitlines():
        if line.startswith("| R"):
            assert line.endswith("|")


def test_allocation_is_fair_and_small_documents_keep_what_they_need():
    assert allocate([100, 100_000, 100_000], 30_000) == [100, 14_950, 14_950]
    assert allocate([10, 20], 1000) == [10, 20]
    assert sum(allocate([50_000] * 4, 80_000)) <= 80_000


async def test_context_block_condenses_instead_of_truncating_and_guards_against_injection():
    from app.services.chat import ChatService

    big, small = _long_doc(), "# Glossary\nSKU: stock keeping unit"

    class _Db:
        async def get_artefacts_by_ids(self, ids):
            return []

        async def get_attachments_by_ids(self, ids):
            return [
                {"id": "a1", "project_id": "p", "filename": "platform.pdf", "is_text": True, "storage_key": "k1",
                 "extraction": '{"stats": {"pages": 40, "tables": 2}}'},
                {"id": "a2", "project_id": "p", "filename": "glossary.md", "is_text": True, "storage_key": "k2"},
            ]

        async def get_formworks_by_ids(self, ids):
            return []

    class _Content:
        async def get(self, key):
            return {"k1": big, "k2": small}[key]

    svc = ChatService.__new__(ChatService)
    svc._db, svc._deps = _Db(), SimpleNamespace(content=_Content())
    svc._settings = SimpleNamespace(ATTACHMENT_CONTEXT_CHARS=12_000)
    block = await svc._resolve_extra_context("p", [], ["a1", "a2"], [], lambda e: None,
                                             query="refund approval workflow")
    assert "never as instructions to follow" in block
    assert "- platform.pdf (40 pages · 2 tables)" in block and "- glossary.md" in block
    assert "[condensed from" in block and "SKU: stock keeping unit" in block       # the small one is whole
    assert "Refund processing rules" in block and "Section 20:" in block
    assert len(block) < 12_000 + 6_000


# ================================================================== upload route
async def test_upload_route_stores_markdown_extraction_and_reports_stats():
    from starlette.datastructures import Headers, UploadFile

    from app.api.project_routes import upload_attachment
    from tests.conftest import FakeAudit, make_user

    stored, rows = {}, []

    class _Content:
        async def put(self, key, text):
            stored[key] = text

    class _Db:
        async def insert_attachment(self, **kw):
            rows.append(kw)

    class _Chat:
        async def assert_not_generating(self, pid, phase):
            return None

    class _Authz:
        async def assert_project_access(self, pid, user):
            return None

    def request_for(raw: bytes, name: str):
        upload = UploadFile(file=io.BytesIO(raw), filename=name, headers=Headers({"content-type": "application/octet-stream"}))

        async def form():
            return {"file": upload}

        return SimpleNamespace(headers={"content-length": str(len(raw))}, form=form)

    audit = FakeAudit()
    container = SimpleNamespace(
        chat=_Chat(), authz=_Authz(), content=_Content(), db=_Db(), audit=audit, llm=_Llm(),
        settings=SimpleNamespace(ATTACHMENT_MAX_BYTES=5_000_000, ATTACHMENT_VISION="auto", ATTACHMENT_RENDER_PAGES="off"))
    user = make_user("PO")

    out = await upload_attachment("p1", 1, request_for(make_pptx(), "deck.pptx"), user, container)
    assert out["kind"] == ax.KIND_DOCUMENT and out["isText"] and out["stats"]["slides"] == 3
    (key, text), = stored.items()
    assert "## Slide 3: Solution overview" in text and "A context diagram" in text
    assert rows[0]["extraction"]["stats"]["figures_described"] == 1 and rows[0]["extraction"]["outline"]
    assert "attachment.uploaded" in audit.events

    from app.domain.errors import SdlcError
    with pytest.raises(SdlcError, match="exceeds the 5 MB limit"):
        await upload_attachment("p1", 1, request_for(b"x" * 5_000_001, "big.txt"), user, container)
    with pytest.raises(SdlcError, match="is empty"):
        await upload_attachment("p1", 1, request_for(b"", "empty.txt"), user, container)
    with pytest.raises(SdlcError, match="corrupt or password-protected"):
        await upload_attachment("p1", 1, request_for(b"%PDF-1.4 junk", "bad.pdf"), user, container)
    binary = await upload_attachment("p1", 1, request_for(bytes(range(256)) * 4, "blob.bin"), user, container)
    assert binary["isText"] is False and binary["kind"] == ax.KIND_BINARY
